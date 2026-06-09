"""
晋辉高考数据分析卡 — Excel 转 JSON 脚本
将山西省高考志愿填报大数据专家版 Excel 转换为前端可用的紧凑 JSON
"""

from __future__ import annotations

import json
import math
import os
import sys
from collections import defaultdict
from typing import Optional

try:
    import openpyxl
except ImportError:
    print("请先安装 openpyxl: pip install openpyxl")
    sys.exit(1)

# ======================== 配置 ========================

EXCEL_PATH = os.path.expanduser(
    "~/Desktop/📊山西高考数据/原数据/山西（2026届）高考志愿填报大数据专家版（持续更新中）.xlsx"
)
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "public", "data")
SHEET_NAME = "山西专家版"

# ======================== 拼音首字母 ========================

# 常用汉字拼音首字母映射 (覆盖高考数据中高频字，避免依赖 pypinyin)
PINYIN_INITIAL_MAP = {
    '北': 'b', '京': 'j', '大': 'd', '学': 'x', '清': 'q', '华': 'h', '理': 'l', '工': 'g',
    '师': 's', '范': 'f', '复': 'f', '旦': 'd', '上': 's', '海': 'h', '交': 'j', '通': 't',
    '浙': 'z', '江': 'j', '南': 'n', '开': 'k', '天': 't', '津': 'j', '武': 'w', '汉': 'h',
    '中': 'z', '山': 's', '西': 'x', '安': 'a', '厦': 'x', '门': 'm', '同': 't', '济': 'j',
    '东': 'd', '科': 'k', '技': 'j', '电': 'd', '子': 'z', '信': 'x', '息': 'x', '数': 's',
    '据': 'j', '软': 'r', '件': 'j', '网': 'w', '络': 'l', '安': 'a', '全': 'q', '人': 'r',
    '民': 'm', '国': 'g', '防': 'f', '经': 'j', '贸': 'm', '财': 'c', '政': 'z', '法': 'f',
    '医': 'y', '药': 'y', '农': 'n', '林': 'l', '语': 'y', '言': 'y', '艺': 'y', '术': 's',
    '体': 't', '育': 'y', '军': 'j', '事': 's', '航': 'h', '空': 'k', '地': 'd', '质': 'z',
    '矿': 'k', '石': 's', '油': 'y', '化': 'h', '建': 'j', '筑': 'z', '水': 's', '力': 'l',
    '新': 'x', '能': 'n', '源': 'y', '材': 'c', '料': 'l', '生': 's', '物': 'w', '环': 'h',
    '境': 'j', '气': 'q', '象': 'x', '海': 'h', '洋': 'y', '统': 't', '计': 'j', '会': 'h',
    '审': 's', '金': 'j', '融': 'r', '管': 'g', '商': 's', '务': 'w', '英': 'y', '文': 'w',
    '日': 'r', '德': 'd', '法': 'f', '俄': 'e', '朝': 'c', '鲜': 'x', '历': 'l', '史': 's',
    '哲': 'z', '社': 's', '公': 'g', '共': 'g', '治': 'z', '马': 'm', '克': 'k', '思': 's',
    '主': 'z', '义': 'y', '心': 'x', '闻': 'w', '传': 'c', '播': 'b', '广': 'g', '告': 'g',
    '设': 's', '计': 'j', '音': 'y', '乐': 'y', '美': 'm', '舞': 'w', '蹈': 'd', '戏': 'x',
    '剧': 'j', '影': 'y', '视': 's', '动': 'd', '画': 'h', '摄': 's', '制': 'z', '临': 'l',
    '床': 'c', '口': 'k', '腔': 'q', '护': 'h', '预': 'y', '基': 'j', '础': 'c', '检': 'j',
    '验': 'y', '麻': 'm', '醉': 'z', '影': 'y', '像': 'x', '康': 'k', '复': 'f', '针': 'z',
    '灸': 'j', '推': 't', '拿': 'n', '藏': 'c', '蒙': 'm', '彝': 'y', '苗': 'm', '壮': 'z',
    '回': 'h', '维': 'w', '哈': 'h', '萨': 's', '塔': 't', '吉': 'j', '乌': 'w', '兹': 'z',
    '别': 'b', '土': 't', '库': 'k', '曼': 'm', '阿': 'a', '塞': 's', '拜': 'b', '疆': 'j',
    '彝': 'y', '藏': 'z', '羌': 'q', '侗': 'd', '瑶': 'y', '傣': 'd', '黎': 'l', '佤': 'w',
    '畲': 's', '高': 'g', '台': 't', '港': 'g', '澳': 'a', '河': 'h', '湖': 'h', '长': 'c',
    '沙': 's', '春': 'c', '成': 'c', '都': 'd', '重': 'c', '庆': 'q', '合': 'h', '肥': 'f',
    '南': 'n', '昌': 'c', '贵': 'g', '阳': 'y', '昆': 'k', '明': 'm', '拉': 'l', '宁': 'n',
    '波': 'b', '温': 'w', '绍': 's', '嘉': 'j', '兴': 'x', '常': 'c', '镇': 'z', '扬': 'y',
    '泰': 't', '盐': 'y', '徐': 'x', '连': 'l', '云': 'y', '淮': 'h', '宿': 's', '迁': 'q',
    '蚌': 'b', '芜': 'w', '淮': 'h', '马': 'm', '鞍': 'a', '铜': 't', '陵': 'l', '安': 'a',
    '宜': 'y', '滁': 'c', '池': 'c', '宣': 'x', '毫': 'h', '阜': 'f', '六': 'l', '锦': 'j',
    '盘': 'p', '葫': 'h', '芦': 'l', '丹': 'd', '辽': 'l', '铁': 't', '抚': 'f', '本': 'b',
    '营': 'y', '阜': 'f', '鄂': 'e', '黄': 'h', '荆': 'j', '孝': 'x', '感': 'g', '咸': 'x',
    '随': 's', '恩': 'e', '施': 's', '岳': 'y', '衡': 'h', '株': 'z', '湘': 'x', '潭': 't',
    '郴': 'c', '永': 'y', '邵': 's', '张': 'z', '家': 'j', '界': 'j', '益': 'y', '娄': 'l',
    '德': 'd', '惠': 'h', '珠': 'z', '汕': 's', '头': 't', '梅': 'm', '湛': 'z', '茂': 'm',
    '肇': 'z', '潮': 'c', '揭': 'j', '韶': 's', '清': 'q', '远': 'y', '江': 'j', '三': 's',
    '桂': 'g', '玉': 'y', '梧': 'w', '钦': 'q', '百': 'b', '色': 's', '贺': 'h', '来': 'l',
    '崇': 'c', '左': 'z', '海': 'h', '口': 'k', '琼': 'q', '三': 's', '亚': 'y', '儋': 'd',
    '柳': 'l', '州': 'z', '莆': 'p', '福': 'f', '泉': 'q', '漳': 'z', '龙': 'l', '岩': 'y',
    '宁': 'n', '夏': 'x', '银': 'y', '川': 'c', '吴': 'w', '忠': 'z', '固': 'g', '原': 'y',
    '兰': 'l', '天': 't', '白': 'b', '银': 'y', '定': 'd', '陇': 'l', '平': 'p', '凉': 'l',
    '庆': 'q', '酒': 'j', '敦': 'd', '煌': 'h', '呼': 'h', '和': 'h', '浩': 'h', '特': 't',
    '包': 'b', '赤': 'c', '峰': 'f', '鄂': 'e', '尔': 'e', '多': 'd', '斯': 's', '通': 't',
    '集': 'j', '巴': 'b', '彦': 'y', '淖': 'n', '尔': 'e', '锡': 'x', '郭': 'g', '勒': 'l',
    '兴': 'x', '延': 'y', '边': 'b', '白': 'b', '城': 'c', '松': 's', '四': 's', '齐': 'q',
    '哈': 'h', '牡': 'm', '佳': 'j', '木': 'm', '大': 'd', '伊': 'y', '鸡': 'j', '鹤': 'h',
    '岗': 'g', '双': 's', '鸭': 'y', '七': 'q', '台': 't', '绥': 's', '黑': 'h', '加': 'j',
    '格': 'g', '达': 'd', '齐': 'q', '曲': 'q', '日': 'r', '喀': 'k', '则': 'z', '昌': 'c',
    '阿': 'a', '里': 'l', '那': 'n', '山': 's', '芝': 'z', '外': 'w', '国': 'g', '创': 'c',
    '业': 'y', '服': 'f', '装': 'z', '纺': 'f', '织': 'z', '食': 's', '品': 'p', '粮': 'l',
    '烟': 'y', '草': 'c', '造': 'z', '纸': 'z', '印': 'y', '刷': 's', '陶': 't', '瓷': 'c',
    '玻': 'b', '璃': 'l', '钢': 'g', '铁': 't', '有': 'y', '色': 's', '金': 'j', '属': 's',
    '机': 'j', '械': 'x', '器': 'q', '仪': 'y', '表': 'b', '汽': 'q', '车': 'c', '船': 'c',
    '舶': 'b', '航': 'h', '宇': 'y', '兵': 'b', '武': 'w', '核': 'h', '农': 'n', '植': 'z',
    '保': 'b', '畜': 'c', '牧': 'm', '兽': 's', '渔': 'y', '草': 'c', '茶': 'c', '园': 'y',
    '艺': 'y', '烟': 'y', '粮': 'l', '储': 'c', '公': 'g', '路': 'l', '铁': 't', '道': 'd',
    '邮': 'y', '快': 'k', '递': 'd', '餐': 'c', '旅': 'l', '酒': 'j', '店': 'd', '会': 'h',
    '展': 'z', '零': 'l', '售': 's', '电': 'd', '商': 's', '市': 's', '场': 'c', '调': 'd',
    '查': 'c', '租': 'z', '赁': 'l', '咨': 'z', '询': 'x', '评': 'p', '估': 'g', '检': 'j',
    '测': 'c', '认': 'r', '证': 'z', '标': 'b', '准': 'z', '化': 'h', '专': 'z', '利': 'l',
    '版': 'b', '权': 'q', '商': 's', '标': 'b', '招': 'z', '标': 'b', '投': 't', '采': 'c',
    '购': 'g', '物': 'w', '业': 'y', '房': 'f', '地': 'd', '产': 'c', '城': 'c', '规': 'g',
    '划': 'h', '景': 'j', '观': 'g', '市': 's', '政': 'z', '照': 'z', '明': 'm', '供': 'g',
    '排': 'p', '燃': 'r', '环': 'h', '卫': 'w', '城': 'c', '管': 'g', '执': 'z', '安': 'a',
    '消': 'x', '防': 'f', '边': 'b', '出': 'c', '入': 'r', '境': 'j', '移': 'y', '民': 'm',
    '普': 'p', '通': 't', '高': 'g', '等': 'd', '院': 'y', '校': 'x', '职': 'z', '培': 'p',
    '训': 'x', '继': 'j', '续': 'x', '育': 'y', '特': 't', '殊': 's', '小': 'x', '幼': 'y',
    '师': 's', '范': 'f', '远': 'y', '程': 'c', '开': 'k', '放': 'f', '广': 'g', '播': 'b',
    '电': 'd', '视': 's', '新': 'x', '媒': 'm', '出': 'c', '编': 'b', '辑': 'j', '记': 'j',
    '者': 'z', '发': 'f', '行': 'x', '印': 'y', '复': 'f', '份': 'f', '文': 'w', '秘': 'm',
    '档': 'd', '案': 'a', '图': 't', '书': 's', '博': 'b', '考': 'k', '古': 'g', '众': 'z',
    '传': 'c', '统': 't', '遗': 'y', '文': 'w', '创': 'c', '族': 'z', '宗': 'z', '民': 'm',
    '间': 'j', '手': 's', '非': 'f', '质': 'z', '保': 'b', '护': 'h', '研': 'y', '究': 'j',
    '自': 'z', '然': 'r', '科': 'k', '工': 'g', '程': 'c', '技': 'j', '实': 's', '试': 's',
    '验': 'y', '班': 'b', '培': 'p', '养': 'y', '拔': 'b', '尖': 'j', '卓': 'z', '越': 'y',
    '院': 'y', '士': 's', '计': 'j', '划': 'h', '强': 'q', '基': 'j', '层': 'c', '选': 'x',
    '修': 'x', '双': 's', '微': 'w', '电': 'd', '集': 'j', '成': 'c', '路': 'l', '半': 'b',
    '导': 'd', '光': 'g', '纤': 'x', '密': 'm', '码': 'm', '感': 'g', '遥': 'y', '探': 't',
    '导': 'd', '航': 'h', '制': 'z', '雷': 'l', '声': 's', '智': 'z', '能': 'n', '算': 's',
    '视': 's', '觉': 'j', '模': 'm', '式': 's', '别': 'b', '对': 'd', '抗': 'k', '虚': 'x',
    '拟': 'n', '现': 'x', '增': 'z', '混': 'h', '量': 'l', '区': 'q', '块': 'k', '隐': 'y',
    '私': 's', '可': 'k', '拓': 't', '展': 'z', '嵌': 'q', '入': 'r', '操': 'c', '作': 'z',
    '移': 'y', '动': 'd', '端': 'd', '桌': 'z', '面': 'm', '服': 'f', '器': 'q', '存': 'c',
    '储': 'c', '云': 'y', '算': 's', '边': 'b', '缘': 'y', '协': 'x', '议': 'y', '接': 'j',
    '口': 'k', '总': 'z', '线': 'x', '编': 'b', '译': 'y', '释': 's', '调': 'd', '优': 'y',
    '并': 'b', '分': 'f', '布': 'b', '容': 'r', '灾': 'z', '备': 'b', '份': 'f', '恢': 'h',
    '升': 's', '级': 'j', '迭': 'd', '代': 'd', '构': 'g', '架': 'j', '组': 'z', '拆': 'c',
}

def get_pinyin_initials(text: str) -> str:
    """获取中文文本的拼音首字母（使用内置映射表）"""
    if not text:
        return ""
    result = []
    for char in text:
        if '一' <= char <= '鿿' or '㐀' <= char <= '䶿':
            initial = PINYIN_INITIAL_MAP.get(char, '')
            result.append(initial if initial else '?')
        elif 'a' <= char.lower() <= 'z':
            result.append(char.lower())
        elif '0' <= char <= '9':
            result.append(char)
    return ''.join(result)


# ======================== 值规范化函数 ========================

def safe_int(value) -> int | None:
    """安全转换为整数"""
    if value is None:
        return None
    try:
        s = str(value).strip().replace(',', '').replace('，', '')
        if s == '' or s == '/' or s == '-' or s == '—' or '待' in s:
            return None
        return int(float(s))
    except (ValueError, TypeError):
        return None

def safe_float(value) -> float | None:
    """安全转换为浮点数"""
    if value is None:
        return None
    try:
        s = str(value).strip().replace('%', '').replace('％', '')
        if s == '' or s == '/' or s == '-' or s == '—':
            return None
        return float(s)
    except (ValueError, TypeError):
        return None

def safe_str(value) -> str:
    """安全转换为字符串"""
    if value is None:
        return ""
    return str(value).strip()

def normalize_baoyan(value) -> float | None:
    """规范化保研率: '12.3%' -> 0.123, '/' -> None"""
    if value is None:
        return None
    s = str(value).strip().replace('%', '').replace('％', '').replace('>', '').replace('≥', '')
    if s == '' or s == '/' or s == '-' or s == '—':
        return None
    try:
        num = float(s)
        if num > 1:
            return num / 100.0
        return num
    except ValueError:
        return None

def parse_xuanke_requirements(xk_str: str) -> list[str]:
    """解析选科要求为数组: '化学和生物' -> ['化学', '生物'], '不限' -> []"""
    xk_str = xk_str.strip()
    if not xk_str or xk_str == '不限':
        return []
    return [s.strip() for s in xk_str.replace('和', ',').replace('、', ',').split(',') if s.strip()]


# ======================== 主转换逻辑 ========================

def convert():
    print(f"📂 读取 Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True, read_only=True)
    ws = wb[SHEET_NAME]
    print(f"   Sheet: {SHEET_NAME}, Rows: {ws.max_row}, Cols: {ws.max_column}")

    # 按科类+批次分组
    groups: dict[str, list[dict]] = defaultdict(list)
    stats = defaultdict(int)
    errors = []

    row_count = 0
    for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        if row_idx % 5000 == 0:
            print(f"   处理中... {row_idx} 行")

        try:
            # 提取关键字段
            year = safe_int(row[0]) if len(row) > 0 else None
            if year != 2025:  # 只处理2025年数据
                continue

            kelei = safe_str(row[2]) if len(row) > 2 else ""
            batch = safe_str(row[3]) if len(row) > 3 else ""

            if not kelei or not batch:
                continue

            # 构造记录
            rec = {
                's': safe_str(row[5]),                    # 院校名称
                'scode': safe_str(row[4]),                # 院校代码
                'g': safe_str(row[7]),                    # 专业组名称
                'gcode': safe_str(row[6]),                # 院校专业组代码
                'm': safe_str(row[10]),                   # 专业全称
                'mn': safe_str(row[11]),                  # 专业名称
                'note': safe_str(row[12]),                # 专业备注
                'level': safe_str(row[13]),               # 专业层次
                'xk': safe_str(row[14]),                  # 选科要求
                'pl': safe_int(row[15]) or 0,             # 计划人数(2025)
                'duration': safe_int(row[16]) or 4,       # 学制
                't': safe_int(row[17]),                   # 学费
                'category': safe_str(row[19]),            # 门类
                'subcategory': safe_str(row[20]),         # 专业类
                'r': safe_int(row[26]),                   # 最低位次(2025)
                'sc': safe_int(row[25]),                  # 最低分(2025)
                'r24': safe_int(row[31]),                 # 最低位次(2024)
                'sc24': safe_int(row[30]),                # 最低分(2024)
                'r23': safe_int(row[36]),                 # 最低位次(2023)
                'sc23': safe_int(row[35]),                # 最低分(2023)
                'p': safe_str(row[43]),                   # 所在省
                'c': safe_str(row[44]),                   # 城市
                'ct': safe_str(row[45]),                  # 城市水平标签
                'tg': safe_str(row[46]),                  # 院校标签
                'schoolLevel': safe_str(row[47]),         # 院校水平
                'affiliation': safe_str(row[49]),         # 隶属单位
                'st': safe_str(row[50]),                  # 类型
                'pp': safe_str(row[51]),                  # 公私性质
                'by': normalize_baoyan(row[53]),          # 保研率
                'schoolRank': safe_int(row[54]),          # 院校排名
                'transfer': safe_str(row[55]) or None,    # 转专业情况
                'xe': safe_str(row[63]) or None,          # 学科评估
                'majorLevel': safe_str(row[64]) or None,  # 专业水平
                'ms': safe_str(row[65]) or None,          # 本专业硕士点
                'ds': safe_str(row[66]) or None,          # 本专业博士点
                'b': batch,                               # 批次
                'kl': kelei,                              # 科类
            }

            # 跳过没有位次数据的行
            if rec['r'] is None and rec['r24'] is None and rec['r23'] is None:
                continue

            # 解析选科要求数组
            rec['xkArr'] = parse_xuanke_requirements(rec['xk'])

            # 计算拼音首字母
            rec['si'] = get_pinyin_initials(rec['s'])
            rec['mi'] = get_pinyin_initials(rec['mn'])

            # 分组键: 科类_批次类型
            if batch in ('本科批', '本科提前批A段', '本科提前批B段', '本科提前批C段'):
                batch_type = 'benke'
            else:
                batch_type = 'zhuanke'

            group_key = f"{'wuli' if kelei == '物理' else 'lishi'}_{batch_type}"
            groups[group_key].append(rec)

            stats['total'] += 1
            stats[f'kelei_{kelei}'] += 1
            stats[f'batch_{batch}'] += 1
            row_count += 1

        except Exception as e:
            errors.append(f"Row {row_idx}: {e}")
            continue

    wb.close()
    print(f"\n✅ 处理完成: {row_count} 条有效记录")

    # 打印统计
    print("\n📊 数据统计:")
    for key in sorted(stats.keys()):
        if key.startswith('kelei_'):
            print(f"   {key}: {stats[key]}")
    for key in sorted(stats.keys()):
        if key.startswith('batch_'):
            print(f"   {key}: {stats[key]}")

    # 输出 JSON 文件
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    print(f"\n💾 输出 JSON 到: {OUTPUT_DIR}")

    for group_key, records in sorted(groups.items()):
        filepath = os.path.join(OUTPUT_DIR, f"{group_key}.json")
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(records, f, ensure_ascii=False, separators=(',', ':'))
        file_size = os.path.getsize(filepath) / (1024 * 1024)
        print(f"   {group_key}.json: {len(records)} 条, {file_size:.1f} MB")

    # 输出省份/城市列表 (供前端筛选器使用)
    provinces = set()
    cities = set()
    school_types = set()
    for recs in groups.values():
        for r in recs:
            if r['p']:
                provinces.add(r['p'])
            if r['c']:
                cities.add(r['c'])
            if r['st']:
                school_types.add(r['st'])

    lookup_data = {
        'provinces': sorted(provinces),
        'cities': sorted(cities),
        'schoolTypes': sorted(school_types),
    }
    lookup_path = os.path.join(OUTPUT_DIR, '_lookup.json')
    with open(lookup_path, 'w', encoding='utf-8') as f:
        json.dump(lookup_data, f, ensure_ascii=False, separators=(',', ':'))
    print(f"   _lookup.json: {len(provinces)} 省, {len(cities)} 市, {len(school_types)} 学校类型")

    if errors:
        print(f"\n⚠️  错误 ({len(errors)} 条):")
        for e in errors[:20]:
            print(f"   {e}")
        if len(errors) > 20:
            print(f"   ... 共 {len(errors)} 条错误")

    print("\n🎉 转换完成!")


if __name__ == '__main__':
    convert()
